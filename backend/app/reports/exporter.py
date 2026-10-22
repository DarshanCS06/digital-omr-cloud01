import io
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any
from sqlalchemy.orm import Session

from backend.app.models.examination import Examination
from backend.app.models.submission import Submission, SubmissionStatus
from backend.app.models.answer_key import AnswerKey
from backend.app.models.batch import UploadedImage, ImageStatus


def sanitize_formula_injection(val: Any) -> Any:
    """Neutralizes spreadsheet formula injection risks by escaping leading trigger characters."""
    if val is None:
        return ""
    s_val = str(val).strip()
    if s_val and s_val[0] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{s_val}"
    return s_val


class ReportExporter:
    @staticmethod
    def generate_csv_results(examination: Examination, db: Session) -> str:
        """Generates sanitized CSV export for examination results."""
        output = io.StringIO()
        writer = csv.writer(output)

        # Active answer key
        active_key = (
            db.query(AnswerKey)
            .filter(AnswerKey.examination_id == examination.id, AnswerKey.is_active == True)
            .order_by(AnswerKey.version.desc())
            .first()
        )
        answer_dict = active_key.answers if active_key else {}

        # Header info
        writer.writerow(["Examination Results Report"])
        writer.writerow(["Title", sanitize_formula_injection(examination.title)])
        writer.writerow(["Code", sanitize_formula_injection(examination.code)])
        writer.writerow(["Subject", sanitize_formula_injection(examination.subject)])
        writer.writerow(["Class", sanitize_formula_injection(examination.class_grade)])
        writer.writerow(["Date", sanitize_formula_injection(examination.exam_date)])
        writer.writerow([])

        # Table headers
        headers = [
            "Submission ID",
            "Roll Number",
            "Student Name",
            "Total Score",
            "Max Score",
            "Percentage",
            "Correct",
            "Incorrect",
            "Blank",
            "Multiple",
            "Needs Review",
            "Status",
        ]
        # Append question columns
        for q in range(1, examination.question_count + 1):
            headers.append(f"Q{q} Answer")
            headers.append(f"Q{q} Correct")

        writer.writerow(headers)

        submissions = (
            db.query(Submission)
            .filter(Submission.examination_id == examination.id)
            .order_by(Submission.id.asc())
            .all()
        )

        for sub in submissions:
            roll = sub.student.roll_number if sub.student else (sub.detected_roll_number or "UNKNOWN")
            name = sub.student.name if sub.student else "Not Assigned"
            
            row = [
                sub.id,
                sanitize_formula_injection(roll),
                sanitize_formula_injection(name),
                sub.total_score,
                sub.max_score,
                f"{sub.percentage}%",
                sub.correct_count,
                sub.incorrect_count,
                sub.blank_count,
                sub.multiple_count,
                sub.review_count,
                sub.status.value,
            ]

            resp_map = {r.question_number: r.final_option for r in sub.responses}
            for q in range(1, examination.question_count + 1):
                ans = resp_map.get(q, "BLANK")
                corr = answer_dict.get(str(q))
                is_correct = "Y" if (ans == corr and corr is not None) else ("N" if ans not in ["BLANK", "REVIEW"] else "-")
                row.append(sanitize_formula_injection(ans))
                row.append(is_correct)

            writer.writerow(row)

        return output.getvalue()

    @staticmethod
    def generate_excel_results(examination: Examination, db: Session) -> bytes:
        """Generates professionally styled Excel (.xlsx) workbook for examination results."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Results Summary"

        # Theme styles
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Dark Blue
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        sub_font = Font(name="Calibri", size=10, italic=True)

        # Title block
        ws["A1"] = f"{examination.title} - Results Evaluation"
        ws["A1"].font = title_font
        ws["A2"] = f"Exam Code: {examination.code} | Class: {examination.class_grade} | Subject: {examination.subject} | Date: {examination.exam_date}"
        ws["A2"].font = sub_font

        # Columns
        headers = [
            "Submission ID", "Roll Number", "Student Name", "Total Score",
            "Max Score", "Percentage", "Correct", "Incorrect", "Blank",
            "Multiple", "Needs Review", "Status"
        ]
        for q in range(1, examination.question_count + 1):
            headers.append(f"Q{q}")

        row_idx = 4
        for col_idx, h_text in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h_text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        submissions = (
            db.query(Submission)
            .filter(Submission.examination_id == examination.id)
            .order_by(Submission.id.asc())
            .all()
        )

        row_idx = 5
        for sub in submissions:
            roll = sub.student.roll_number if sub.student else (sub.detected_roll_number or "UNKNOWN")
            name = sub.student.name if sub.student else "Not Assigned"

            row_data = [
                sub.id,
                sanitize_formula_injection(roll),
                sanitize_formula_injection(name),
                sub.total_score,
                sub.max_score,
                f"{sub.percentage}%",
                sub.correct_count,
                sub.incorrect_count,
                sub.blank_count,
                sub.multiple_count,
                sub.review_count,
                sub.status.value,
            ]

            resp_map = {r.question_number: r.final_option for r in sub.responses}
            for q in range(1, examination.question_count + 1):
                row_data.append(sanitize_formula_injection(resp_map.get(q, "BLANK")))

            for col_idx, val in enumerate(row_data, 1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.alignment = Alignment(horizontal="center" if col_idx != 3 else "left")

            row_idx += 1

        # Second Sheet: Sheet-wise Error Report
        ws_errors = wb.create_sheet(title="Upload & Error Report")
        err_headers = ["Image ID", "Filename", "File Size (KB)", "Status", "Error Message", "Created At"]
        for col_idx, h_text in enumerate(err_headers, 1):
            cell = ws_errors.cell(row=1, column=col_idx, value=h_text)
            cell.fill = header_fill
            cell.font = header_font

        images = (
            db.query(UploadedImage)
            .filter(UploadedImage.examination_id == examination.id)
            .order_by(UploadedImage.id.asc())
            .all()
        )

        err_row = 2
        for img in images:
            ws_errors.cell(row=err_row, column=1, value=img.id)
            ws_errors.cell(row=err_row, column=2, value=sanitize_formula_injection(img.original_filename))
            ws_errors.cell(row=err_row, column=3, value=round(img.file_size_bytes / 1024.0, 1))
            ws_errors.cell(row=err_row, column=4, value=img.status.value)
            ws_errors.cell(row=err_row, column=5, value=sanitize_formula_injection(img.error_message or "None"))
            ws_errors.cell(row=err_row, column=6, value=img.created_at.strftime("%Y-%m-%d %H:%M:%S"))
            err_row += 1

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_pdf_results(examination: Examination, db: Session, batch_id: int = None) -> bytes:
        """Generates a printable A4 PDF Results Report for the examination or specific batch."""
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.lib import colors

        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=A4)
        page_w, page_h = A4

        sub_q = db.query(Submission).filter(Submission.examination_id == examination.id)
        img_q = db.query(UploadedImage).filter(UploadedImage.examination_id == examination.id)
        if batch_id is not None:
            img_ids = [
                r.id
                for r in db.query(UploadedImage.id).filter(
                    UploadedImage.examination_id == examination.id,
                    UploadedImage.batch_id == batch_id,
                ).all()
            ]
            sub_q = sub_q.filter(Submission.uploaded_image_id.in_(img_ids))
            img_q = img_q.filter(UploadedImage.batch_id == batch_id)

        submissions = sub_q.order_by(Submission.id.asc()).all()
        images = img_q.order_by(UploadedImage.id.asc()).all()

        total_pages = len(images) if images else len(submissions)
        eval_count = len(submissions)
        review_count = sum(1 for s in submissions if s.status == SubmissionStatus.NEEDS_REVIEW)
        failed_imgs = [im for im in images if im.status == ImageStatus.FAILED]
        dup_count = sum(1 for s in submissions if s.is_duplicate)

        def _draw_header(page_num: int):
            c.setFillColor(colors.HexColor("#1E3A8A"))
            c.rect(30, page_h - 80, page_w - 60, 50, fill=1, stroke=0)
            c.setFillColor(colors.white)
            c.setFont("Helvetica-Bold", 14)
            c.drawString(44, page_h - 52, "VIDYAKASHI STUDY CIRCLE — OMR EVALUATION REPORT")
            c.setFont("Helvetica", 9)
            c.drawString(
                44,
                page_h - 68,
                f"Exam: {examination.title} ({examination.code}) | Subject: {examination.subject} | Class: {examination.class_grade} | Page {page_num}",
            )

        page_num = 1
        _draw_header(page_num)

        # Summary KPI row
        y = page_h - 105
        c.setFillColor(colors.HexColor("#F1F5F9"))
        c.rect(30, y - 22, page_w - 60, 28, fill=1, stroke=1)
        c.setFillColor(colors.HexColor("#0F172A"))
        c.setFont("Helvetica-Bold", 9)
        summary_line = (
            f"Total Pages: {total_pages}   |   Evaluated: {eval_count}   |   "
            f"Needs Review: {review_count}   |   Failed: {len(failed_imgs)}   |   "
            f"Duplicate Roll Nos: {dup_count}"
        )
        c.drawString(40, y - 11, summary_line)

        # Table Header
        y -= 42
        col_x = [34, 70, 145, 275, 320, 365, 410, 465, 515]
        col_headers = ["#", "Roll No.", "Student Name", "Right", "Wrong", "Blank", "Marks", "%", "Status"]
        c.setFillColor(colors.HexColor("#1E293B"))
        c.rect(30, y - 6, page_w - 60, 20, fill=1, stroke=0)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 8.5)
        for cx, h_txt in zip(col_x, col_headers):
            c.drawString(cx, y, h_txt)

        y -= 20
        c.setFont("Helvetica", 8.5)
        for idx, sub in enumerate(submissions, start=1):
            if y < 70:
                c.showPage()
                page_num += 1
                _draw_header(page_num)
                y = page_h - 105
                c.setFont("Helvetica", 8.5)

            roll = sub.student.roll_number if sub.student else (sub.detected_roll_number or "N/A")
            name = sub.student.name if sub.student else "Unassigned"
            status_str = "DUPLICATE" if sub.is_duplicate else sub.status.value

            if idx % 2 == 0:
                c.setFillColor(colors.HexColor("#F8FAFC"))
                c.rect(30, y - 5, page_w - 60, 17, fill=1, stroke=0)

            c.setFillColor(colors.HexColor("#0F172A"))
            c.drawString(col_x[0], y, str(idx))
            c.drawString(col_x[1], y, str(roll)[:12])
            c.drawString(col_x[2], y, str(name)[:22])
            c.drawString(col_x[3], y, str(sub.correct_count))
            c.drawString(col_x[4], y, str(sub.incorrect_count))
            c.drawString(col_x[5], y, str(sub.blank_count))
            c.drawString(col_x[6], y, f"{sub.total_score}/{sub.max_score}")
            c.drawString(col_x[7], y, f"{sub.percentage}%")
            c.drawString(col_x[8], y, status_str[:11])
            y -= 17

        if failed_imgs:
            y -= 15
            if y < 100:
                c.showPage()
                page_num += 1
                _draw_header(page_num)
                y = page_h - 105
            c.setFillColor(colors.HexColor("#991B1B"))
            c.setFont("Helvetica-Bold", 9.5)
            c.drawString(34, y, f"Failed / Rejected Pages ({len(failed_imgs)}):")
            y -= 16
            c.setFont("Helvetica", 8)
            for fim in failed_imgs:
                if y < 60:
                    c.showPage()
                    page_num += 1
                    _draw_header(page_num)
                    y = page_h - 105
                    c.setFont("Helvetica", 8)
                c.drawString(38, y, f"• {fim.original_filename}: {fim.error_message or 'Failed validation'}")
                y -= 14

        c.save()
        return buf.getvalue()
